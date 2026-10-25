"""Turn Case_Study_Report.md into a styled Word document. Usage (from case_study/): python scripts/build_report_docx.py"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DEEP, OCEAN, MUTED = RGBColor(0x06, 0x32, 0x3F), RGBColor(0x0E, 0x6E, 0x80), RGBColor(0x4F, 0x6B, 0x73)
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|\[[^\]]+\]\([^)]+\)|https?://[^\s)]+)")


def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear"); el.set(qn("w:color"), "auto"); el.set(qn("w:fill"), fill)
    pr.append(el)


def borders(table, color="D5E6E4", size=4):
    pr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(size)); e.set(qn("w:color"), color)
        b.append(e)
    pr.append(b)
    m = OxmlElement("w:tblCellMar")
    for edge, w in (("top", 50), ("bottom", 50), ("left", 90), ("right", 90)):
        e = OxmlElement(f"w:{edge}"); e.set(qn("w:w"), str(w)); e.set(qn("w:type"), "dxa"); m.append(e)
    pr.append(m)


def hyperlink(par, url, text, size=None):
    rid = par.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
    r = OxmlElement("w:r"); rp = OxmlElement("w:rPr")
    for tag, val in (("w:color", "0E6E80"), ("w:u", "single")):
        e = OxmlElement(tag); e.set(qn("w:val"), val); rp.append(e)
    if size:
        sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(size * 2))); rp.append(sz)
    r.append(rp); t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve"); r.append(t); h.append(r)
    par._p.append(h)


def inline(par, text, size=None, color=None):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = par.add_run(part[1:-1]); r.italic = True
        elif part.startswith("["):
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            hyperlink(par, m.group(2), m.group(1), size); continue
        elif part.startswith("http"):
            url = part.rstrip(".,;")
            hyperlink(par, url, url, size)
            if part != url:
                par.add_run(part[len(url):])
            continue
        else:
            r = par.add_run(part)
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = color


def para(doc, text, size=None, align=None, after=6, before=0, color=None, indent=None, hanging=None, keep=False):
    p = doc.add_paragraph()
    inline(p, text, size, color)
    f = p.paragraph_format
    f.space_after, f.space_before, f.line_spacing = Pt(after), Pt(before), 1.12
    if align: p.alignment = align
    if indent is not None: f.left_indent = Cm(indent)
    if hanging is not None: f.first_line_indent = Cm(-hanging)
    if keep: f.keep_with_next = True
    return p


def heading(doc, text, level):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = True; r.font.name = "Georgia"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Georgia")
    r.font.size, r.font.color.rgb = (Pt(15), OCEAN) if level == 1 else (Pt(12), DEEP)
    f = p.paragraph_format
    f.space_before, f.space_after, f.keep_with_next = Pt(16 if level == 1 else 10), Pt(5), True
    return p


def add_table(doc, rows):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j); c.text = ""
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            inline(p, val, 9, RGBColor(255, 255, 255) if i == 0 else None)
            for r in p.runs:
                r.font.size = Pt(9)
                if i == 0: r.bold = True; r.font.color.rgb = RGBColor(255, 255, 255)
            if i == 0: shade(c, "0E6E80")
            elif i % 2 == 0: shade(c, "EEF8F7")
    hdr = t.rows[0]._tr.get_or_add_trPr(); e = OxmlElement("w:tblHeader"); e.set(qn("w:val"), "true"); hdr.append(e)
    for r in t.rows:
        cant = OxmlElement("w:cantSplit"); cant.set(qn("w:val"), "true"); r._tr.get_or_add_trPr().append(cant)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_code(doc, lines):
    t = doc.add_table(rows=1, cols=1); borders(t, "D5E6E4")
    c = t.cell(0, 0); shade(c, "F3F7F7"); c.text = ""
    p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    r = p.add_run("\n".join(lines)); r.font.name = "Consolas"; r.font.size = Pt(8)
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def footer_numbers(section, left):
    p = section.footer.paragraphs[0]
    p.text = ""
    r = p.add_run(left + "   ·   page "); r.font.size = Pt(8.5); r.font.color.rgb = MUTED
    run = p.add_run(); run.font.size = Pt(8.5); run.font.color.rgb = MUTED
    for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), kind)
        else:
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt
        run._r.append(e)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def main():
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "Case_Study_Report.md"
    lines = src.read_text(encoding="utf-8").split("\n")
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin"): setattr(sec, side, Cm(2.3))
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.2)
    st = doc.styles["Normal"]; st.font.name, st.font.size = "Calibri", Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    footer_numbers(sec, "Devang Pradeep Mishra  ·  CSE4143")

    title = lines[0][2:]
    author = next(l for l in lines[1:] if l.strip())
    link_line = next(l for l in lines if l.startswith("Code and results"))
    box = doc.add_table(rows=1, cols=1); c = box.cell(0, 0); shade(c, "06323F"); c.text = ""
    p = c.paragraphs[0]; r = p.add_run("CSE4143 RECOMMENDER SYSTEM  ·  CASE STUDY" if "Reader" not in title else "CSE4143 RECOMMENDER SYSTEM  ·  COMPANION GUIDE"); r.font.size = Pt(8.5); r.bold = True
    r.font.color.rgb = RGBColor(0x2E, 0xC4, 0xB6); p.paragraph_format.space_before = Pt(10)
    p = c.add_paragraph(); r = p.add_run(title); r.font.name, r.font.size, r.bold = "Georgia", Pt(21), True
    r.font.color.rgb = RGBColor(255, 255, 255); p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(6), Pt(8)
    p = c.add_paragraph(); r = p.add_run(author); r.font.size = Pt(9.5); r.font.color.rgb = RGBColor(0xBF, 0xE4, 0xE6)
    p = c.add_paragraph(); p.paragraph_format.space_after = Pt(10)
    inline(p, link_line, 9.5, RGBColor(0xBF, 0xE4, 0xE6))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    i, fig_no = 0, 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("# ") or ln == author or ln == link_line or not ln.strip():
            i += 1; continue
        if ln.startswith("```"):
            block = []; i += 1
            while not lines[i].startswith("```"): block.append(lines[i]); i += 1
            add_code(doc, block); i += 1; continue
        if ln.startswith("## "):
            heading(doc, ln[3:], 1)
        elif ln.startswith("### "):
            heading(doc, ln[4:], 2)
        elif ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [x.strip() for x in lines[i].strip().strip("|").split("|")]
                if not all(set(x) <= set("- ") for x in cells): rows.append(cells)
                i += 1
            add_table(doc, rows); continue
        elif ln.startswith("!["):
            m = re.match(r"!\[(.*)\]\((.*)\)", ln)
            pic = doc.add_paragraph(); pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pic.paragraph_format.keep_with_next, pic.paragraph_format.space_after = True, Pt(2)
            pic.add_run().add_picture(str(ROOT / m.group(2)), width=Cm(14.2))
            para(doc, m.group(1), 8.5, WD_ALIGN_PARAGRAPH.CENTER, after=10, color=MUTED)
        elif ln.startswith("- "):
            p = para(doc, "•\t" + ln[2:], after=3, indent=0.9, hanging=0.5)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(0.9))
        elif re.match(r"\d+\. ", ln):
            para(doc, ln, 9.5, after=3, indent=0.8, hanging=0.8)
        else:
            para(doc, ln)
        i += 1
    out = src.with_suffix(".docx")
    doc.save(out)
    print("saved", out)


if __name__ == "__main__":
    sys.exit(main())
